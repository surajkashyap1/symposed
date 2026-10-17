"""The editable files that come with the guide.

  tracker.xlsx         every suggested paper, ready for screening (decision drop
                       downs, exclusion reasons from the project's decision tree),
                       plus Data extraction, Risk of bias (the chosen tool's
                       domains and judgements) and Search log sheets
  protocol.docx        the draft protocol (SOP), for the team to edit
  prospero_draft.docx  the PROSPERO fields filled in (systematic reviews only)

All three are built from the same guide model as the PDF, so names, numbers and
step titles match.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CLARET, INK, TEAL, SLATE = "7A2E2E", "1B2A2E", "0F6F66", "2F5D7C"
DECISIONS = ("Include", "Exclude", "Further review needed")
YES_NO = ("Yes", "No")

_THIN = Side(style="thin", color="E5DFD3")
_WRAP = Alignment(wrap_text=True, vertical="top")


# -- tracker spreadsheet ------------------------------------------------------------
def _header(ws, columns: list[tuple[str, float]], colour: str, row: int = 1) -> None:
    for i, (name, width) in enumerate(columns, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=colour)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=_THIN)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=3)


def _dropdown(ws, wb_lists, name: str, options: list[str], col: int, first: int,
              last: int) -> DataValidation:
    """A list validation fed from the hidden Lists sheet (no 255 character limit)."""
    lists = wb_lists
    col_idx = lists.max_column + 1 if lists.max_column > 1 or lists["A1"].value else 1
    lists.cell(row=1, column=col_idx, value=name).font = Font(bold=True)
    for i, opt in enumerate(options, 2):
        lists.cell(row=i, column=col_idx, value=opt)
    letter = get_column_letter(col_idx)
    dv = DataValidation(type="list", allow_blank=True,
                        formula1=f"=Lists!${letter}$2:${letter}${len(options) + 1}")
    dv.error, dv.errorTitle = "Pick an option from the list, or add one on the Lists sheet.", name
    ws.add_data_validation(dv)
    col_letter = get_column_letter(col)
    dv.add(f"{col_letter}{first}:{col_letter}{last}")
    return dv


def write_tracker(m: dict[str, Any], path: str | Path, *, blank_rows: int = 300) -> Path:
    wb = Workbook()
    intro = wb.active
    intro.title = "How to use"
    lists = wb.create_sheet("Lists")
    screening = wb.create_sheet("Screening")
    extraction = wb.create_sheet("Data extraction")
    bias = wb.create_sheet("Risk of bias")
    log = wb.create_sheet("Search log")

    # How to use
    intro.column_dimensions["A"].width = 110
    lines = [
        (m["title"], Font(bold=True, size=14, color=CLARET)),
        ("Paper tracker made for your project by Symposed. Every sheet is a starting point: "
         "add, remove or rename columns as you see fit.", Font(italic=True)),
        ("", None),
        ("Screening: every paper we suggest is already listed (marked 'Symposed suggestion'). "
         "Add each paper your own searches find underneath. Each screener picks a decision "
         "from the drop down: Include, Exclude or Further review needed. For exclusions, "
         "pick a reason. The reasons come from the decision tree in step 4 of your guide.", None),
        ("Data extraction: one row per included paper. The columns are the details step 5 "
         "of your guide asks you to extract.", None),
        (f"Risk of bias: one row per included study, one column per part of "
         f"{', '.join(t['name'] for t in m['rob_tools']) or 'your tool'}, as in step 6.", None),
        ("Search log: record every search you run (database, date, the exact search and the "
         "number of results). You need this for your methods and PRISMA diagram.", None),
        ("Lists: the options behind each drop down. Edit them here to change the choices.", None),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = intro.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if font:
            c.font = font

    # Screening
    cols = [("Row", 6), ("PMID or ID", 13), ("DOI", 22), ("Year first published", 10),
            ("First author", 16), ("Title", 48), ("Abstract", 60), ("Journal", 22),
            ("Likely publication type", 20), ("Type from", 12), ("Link", 30),
            ("Source", 18), ("Screener 1 decision", 16), ("Screener 2 decision", 16),
            ("Agreed decision", 16), ("Reason for exclusion", 30),
            ("Full text obtained?", 12), ("Notes", 40)]
    _header(screening, cols, INK)
    for r, p in enumerate(m["papers"], 2):
        values = [r - 1, p["pmid"], p.get("doi", ""), p.get("first_published") or p.get("year"),
                  p["first_author"], p["title"], p.get("abstract", ""), p["journal"],
                  p["publication_type"], p["publication_type_basis"], p["url"],
                  "Symposed suggestion"]
        for c, v in enumerate(values, 1):
            cell = screening.cell(row=r, column=c, value=v)
            cell.alignment = _WRAP
        screening.cell(row=r, column=11).hyperlink = p["url"]
        screening.cell(row=r, column=11).font = Font(color=CLARET, underline="single")
        screening.row_dimensions[r].height = 60
    last = len(m["papers"]) + 1 + blank_rows
    for c in (13, 14, 15):
        _dropdown(screening, lists, ["Screener 1", "Screener 2", "Agreed"][c - 13] + " decision",
                  list(DECISIONS), c, 2, last)
    _dropdown(screening, lists, "Reason for exclusion", m.get("exclusion_reasons")
              or ["Other (write in notes)"], 16, 2, last)
    _dropdown(screening, lists, "Full text obtained", list(YES_NO), 17, 2, last)
    screening.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    _shade_decisions(screening, last)

    # Data extraction
    ex_cols = [("Row", 6), ("PMID or ID", 13)] + [(f, 26) for f in m["extraction_fields"]]
    ex_cols.append(("Extracted by", 14))
    ex_cols.append(("Checked by", 14))
    _header(extraction, ex_cols, TEAL)

    # Risk of bias, one block of columns per tool
    rb_cols: list[tuple[str, float]] = [("Row", 6), ("PMID or ID", 13), ("Study", 22)]
    spans: list[tuple[int, list[str]]] = []
    for t in m["rob_tools"]:
        start = len(rb_cols) + 1
        names = [f"{t['name']}: {d}" for d in t["domains"]] + [f"{t['name']}: overall"]
        rb_cols += [(n, 20) for n in names]
        spans.append((start, list(t["judgements"])))
    rb_cols += [("Assessed by", 14), ("Notes", 40)]
    _header(bias, rb_cols, SLATE)
    for (start, judgements), t in zip(spans, m["rob_tools"]):
        n = len(t["domains"]) + 1
        dv = _dropdown(bias, lists, f"{t['name']} judgement", judgements, start, 2, 200)
        for c in range(start + 1, start + n):
            dv.add(f"{get_column_letter(c)}2:{get_column_letter(c)}200")

    # Search log, pre filled with the searches in the guide
    _header(log, [("Search strategy", 12), ("Database", 18), ("Date run", 14),
                  ("Exact search", 90), ("Results", 10), ("Run by", 14), ("Notes", 30)], CLARET)
    for r, s in enumerate(m["searches"], 2):
        for c, v in enumerate([s["number"], "PubMed", m["search_date"], s["query"],
                               s.get("count"), "Symposed", tidy_note(s)], 1):
            log.cell(row=r, column=c, value=v).alignment = _WRAP

    lists.sheet_state = "hidden"
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    path = Path(path)
    wb.save(str(path))
    return path


def tidy_note(s: dict) -> str:
    if s.get("found_papers"):
        return "Found the suggested papers in your guide"
    return "Recommended in your guide" if s.get("recommended") else ""


def _shade_decisions(ws, last: int) -> None:
    from openpyxl.formatting.rule import FormulaRule

    for value, colour in (("Include", "DDF0EC"), ("Exclude", "FBE4E1"),
                          ("Further review needed", "FBF0D9")):
        for col in ("M", "N", "O"):
            ws.conditional_formatting.add(
                f"{col}2:{col}{last}",
                FormulaRule(formula=[f'{col}2="{value}"'],
                            fill=PatternFill("solid", fgColor=colour)))


# -- Word documents ----------------------------------------------------------------
def _doc(title: str, subtitle: str):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        doc.styles[name].font.color.rgb = RGBColor.from_string(CLARET if name != "Heading 2"
                                                               else INK)
        doc.styles[name].font.size = Pt(size)
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.runs[0].italic = True
    return doc


def _note(doc, text: str, fill: str = "FBF0D9") -> None:
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.text = text
    _shade(cell, fill)
    doc.add_paragraph()


def _shade(cell, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _table(doc, header: list[str], rows: list[list[str]], colour: str = INK) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = text
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        _shade(cell, colour)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = str(text)
    doc.add_paragraph()


EDIT_NOTE = ("This is a starting point drafted by Symposed. Edit it as you see fit with your "
             "team and supervisor. Delete this note when you are done.")


def write_protocol(m: dict[str, Any], path: str | Path) -> Path:
    doc = _doc(f"Protocol: {m['title']}",
               f"{m['publication_type'].capitalize()} · draft prepared {m['generated_on']}")
    _note(doc, EDIT_NOTE)
    doc.add_heading("1. Review team", level=1)
    doc.add_paragraph("Names, roles and affiliations: [add your team]. Supervisor: [add if you "
                      "have one].")
    doc.add_heading("2. Background and rationale", level=1)
    doc.add_paragraph(m["rationale"])
    doc.add_heading("3. Aim and objectives", level=1)
    doc.add_paragraph(m["aims"])
    for o in m["objectives"]:
        doc.add_paragraph(o, style="List Bullet")
    doc.add_heading("4. Eligibility criteria", level=1)
    c = m["criteria"]
    rows = [["Population", c["population"]], ["Intervention or exposure", c["intervention"]],
            ["Comparator", c["comparator"]], ["Outcomes", c["outcomes"]],
            ["Study designs", c["designs"]]]
    if c.get("setting"):
        rows.append(["Setting", c["setting"]])
    rows += [["Also required", x] for x in c["other_inclusion"]]
    rows += [["Exclude", x] for x in c["exclusion"]]
    _table(doc, ["Criterion", "Rule"], rows, TEAL)
    if m["criteria_reasoning"]:
        doc.add_heading("Reasoning", level=2)
        for r in m["criteria_reasoning"]:
            p = doc.add_paragraph()
            p.add_run(f"{r['criterion']}: ").bold = True
            p.add_run(r["reasoning"])
    doc.add_heading("5. Information sources and search", level=1)
    doc.add_paragraph("Databases: PubMed (MEDLINE)" + (", " + ", ".join(
        d for d in m["profile"]["databases"] if d not in ("PubMed only", "Unsure"))
        if m["profile"]["databases"] else "") + ". Grey literature: trial registries, "
        "conference abstracts and reference lists.")
    rec = next((s for s in m["searches"] if s["recommended"]), m["searches"][0] if m["searches"]
               else None)
    if rec:
        doc.add_paragraph(f"PubMed search strategy (search strategy {rec['number']} in the "
                          f"guide, {rec['count'] if rec['count'] is not None else 'n/a'} records "
                          f"on {m['search_date']}):")
        q = doc.add_paragraph(rec["query"])
        q.runs[0].font.name = "Consolas"
        q.runs[0].font.size = Pt(9)
    doc.add_heading("6. Study selection", level=1)
    doc.add_paragraph("Two reviewers will screen titles and abstracts independently in Rayyan "
                      "with blinding switched on, then the full texts of potentially eligible "
                      "papers. Disagreements will be resolved by discussion, or by a third "
                      "reviewer.")
    if m["decision_tree"]:
        doc.add_paragraph("Screening questions, asked in order (a no excludes the paper):")
        for i, q in enumerate(m["decision_tree"], 1):
            doc.add_paragraph(f"{i}. {q['question']} If no: {q['no_label']}.")
    doc.add_heading("7. Data extraction", level=1)
    doc.add_paragraph("One reviewer will extract data using a piloted form and a second will "
                      "check it. Items:")
    for f in m["extraction_fields"]:
        doc.add_paragraph(f, style="List Bullet")
    doc.add_heading("8. Risk of bias", level=1)
    for t in m["rob_tools"]:
        doc.add_paragraph(f"{t['name']} ({t['full_name']}) for {t['use_for']}. Domains: "
                          f"{', '.join(t['domains'])}. Two reviewers will judge each study "
                          f"independently.")
    doc.add_heading("9. Data synthesis", level=1)
    doc.add_paragraph(m["analysis"])
    doc.add_heading("10. Reporting", level=1)
    for r in m["reporting"]:
        doc.add_paragraph(r, style="List Bullet")
    doc.add_heading("11. Timeline", level=1)
    _table(doc, ["Step", "Hours each", "Weeks"],
           [[f"{t['number']}. {t['name']}", f"{t['hours']:g}",
             f"{t['start_week']} to {t['end_week']}"] for t in m["timeline"]], SLATE)
    path = Path(path)
    doc.save(str(path))
    return path


PROSPERO_LABELS = {
    "review_title": "Review title", "review_question": "Review question",
    "searches": "Searches", "condition_or_domain": "Condition or domain being studied",
    "population": "Participants or population",
    "intervention_exposure": "Intervention(s) or exposure(s)",
    "comparator_control": "Comparator(s) or control(s)", "types_of_study": "Types of study to be included",
    "main_outcomes": "Main outcome(s)", "additional_outcomes": "Additional outcome(s)",
    "data_extraction": "Data extraction (selection and coding)",
    "risk_of_bias_assessment": "Risk of bias (quality) assessment",
    "strategy_for_synthesis": "Strategy for data synthesis",
    "subgroup_analyses": "Analysis of subgroups or subsets",
}
YOU_FILL = ("Named contact and email", "Review team members and their affiliations",
            "Anticipated start and completion dates", "Funding sources",
            "Conflicts of interest", "Stage of the review when you register",
            "Country and language")


def write_prospero(m: dict[str, Any], path: str | Path) -> Path:
    doc = _doc("PROSPERO registration draft", m["title"])
    _note(doc, "Copy each field into the PROSPERO form at https://www.crd.york.ac.uk/prospero/. "
               "These drafts come from your protocol and are a starting point. Edit them "
               "before you submit, and register before you start screening.")
    _table(doc, ["PROSPERO field", "Draft for you to edit"],
           [[PROSPERO_LABELS.get(k, k.replace("_", " ").capitalize()), v]
            for k, v in (m.get("prospero") or {}).items()], CLARET)
    doc.add_heading("Fields only you can fill in", level=1)
    for f in YOU_FILL:
        doc.add_paragraph(f, style="List Bullet")
    path = Path(path)
    doc.save(str(path))
    return path


def write_all(m: dict[str, Any], directory: str | Path) -> dict[str, Path]:
    d = Path(directory)
    out = {"tracker": write_tracker(m, d / "tracker.xlsx"),
           "protocol": write_protocol(m, d / "protocol.docx")}
    if m.get("systematic") and m.get("prospero"):
        out["prospero"] = write_prospero(m, d / "prospero_draft.docx")
    return out
