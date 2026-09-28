"""Render the Publication Guide as an editable Word document (spec §6).

Every section comes from `GuideContent`, assembled from retrieved records and
checked model output. Sections follow the spec's order; systematic reviews get
all of them, lighter review types get the ones that apply (no PROSPERO entry or
PRISMA diagram, a synthesis approach and a journal shortlist instead). The novelty
statement, registration urgency and disclaimers are verbatim.
"""

from __future__ import annotations

from typing import Optional

from docx import Document
from docx.shared import Pt

from .guide_content import (
    ACKNOWLEDGEMENT_WORDING,
    AUTHOR_AI_NOTE,
    COLLABORATION_PROMPT,
    DISCLAIMERS,
    GREY_LITERATURE,
    NOVELTY_STATEMENT,
    REGISTRATION_URGENCY,
    SCREENING_ESTIMATE_NOTE,
    SCREENING_PLATFORM,
    GuideContent,
    rejection_reasons,
)
from .landscape import Landscape
from .screening import NOT_STATED, STATUSES
from .sources.pubmed import Paper
from .workspace import Workspace


def _format_citation(paper: Paper) -> str:
    authors = list(paper.authors)
    if not authors:
        who = "[No author listed]"
    elif len(authors) > 3:
        who = ", ".join(authors[:3]) + ", et al."
    else:
        who = ", ".join(authors)
    year = f" ({paper.year})" if paper.year else ""
    journal = f" {paper.journal}." if paper.journal else ""
    label = "Europe PMC" if paper.is_preprint else "PMID:"
    ids = f" {label} {paper.pmid}." if paper.pmid else ""
    if paper.doi:
        ids += f" DOI: {paper.doi}."
    return f"{who}{year}. {paper.title}.{journal}{ids}".strip()


def _truncate(text: str, n: int = 90) -> str:
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _table(doc, header: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = str(text)
    return table


def _bullets(doc, items) -> None:
    for item in items:
        doc.add_paragraph(str(item), style="List Bullet")


def build_guide(
    workspace: Workspace,
    content: GuideContent,
    *,
    landscape: Optional[Landscape] = None,
) -> str:
    """Write the guide to `workspace.guide_path` and return that path."""
    c = content
    doc = Document()
    n = 0

    def section(title: str) -> None:
        nonlocal n
        n += 1
        doc.add_heading(f"{n}. {title}", level=1)

    # 1. Title page
    doc.add_heading(c.title, level=0)
    meta = doc.add_paragraph()
    meta.add_run("Publication type: ").bold = True
    meta.add_run(f"{c.publication_type.capitalize()}\n")
    meta.add_run("Guide generated: ").bold = True
    meta.add_run(f"{c.generated_on}\n")
    meta.add_run("Searches run: ").bold = True
    meta.add_run(c.search_date)

    # 2. Novelty statement (verbatim, near the front)
    section("Novelty statement")
    doc.add_paragraph(NOVELTY_STATEMENT.format(date=c.novelty_date))

    # 3. The research question
    section("The research question (PICO)")
    _table(doc, ["Element", "Definition"], [[k, v] for k, v in c.pico])

    # 4. Rationale
    section("Rationale")
    doc.add_paragraph(c.rationale)
    if c.references:
        doc.add_paragraph("References", style="Heading 3")
        for i, ref in enumerate(c.references, 1):
            doc.add_paragraph(f"[{i}] {ref}")

    # 5. Similar existing publications
    section("Similar existing work and how this question differs")
    if not c.similar:
        doc.add_paragraph("No closely similar reviews or registered protocols were found.")
    for i, s in enumerate(c.similar, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(f"{s['title']} ").bold = True
        p.add_run(f"({s['url']}; similarity {s['score']:.2f})")
        if c.differences.get(i):
            doc.add_paragraph(f"How yours differs: {c.differences[i]}")

    # 6. Feasibility
    section("Feasibility")
    f = c.feasibility
    _bullets(doc, [
        f"Eligible studies: about {f['eligible_screened']} plausibly eligible in our "
        f"preliminary screen ({f['eligible_triage']} by count before screening).",
        f"Expected screening volume: about {f['records_to_screen']} records from PubMed "
        "alone; other databases will add more.",
        f"Recommended team: at least {f['recommended_team']} people, so every record is "
        "screened independently by two.",
    ])

    # 7. Search strategies
    section("Search strategies")
    s = c.strategies
    for i, st in enumerate(s["strategies"]):
        rec = " (recommended)" if i == s["recommended"] else ""
        doc.add_paragraph(f"{st.name}{rec}", style="Heading 3")
        doc.add_paragraph(st.purpose)
        doc.add_paragraph(st.pubmed_query, style="Intense Quote")
        count = "not available" if st.count is None else f"{st.count:,}"
        doc.add_paragraph(f"PubMed records on {c.search_date}: {count}")
    if s.get("recommendation_reason"):
        doc.add_paragraph(f"Why we recommend it: {s['recommendation_reason']}")

    # 8. Translated syntax
    section("Translated syntax for databases you search yourself")
    doc.add_paragraph("These translations were not run by us (we have no access). "
                      "Run and adjust them in each database.")
    doc.add_paragraph("Embase (Emtree)", style="Heading 3")
    doc.add_paragraph(s.get("embase_emtree", ""), style="Intense Quote")
    doc.add_paragraph("Cochrane CENTRAL", style="Heading 3")
    doc.add_paragraph(s.get("cochrane_central", ""), style="Intense Quote")

    # 9. Suggested papers table
    section("Suggested papers")
    doc.add_paragraph(SCREENING_ESTIMATE_NOTE)
    rows = []
    for status in STATUSES:
        for sp in c.papers.by_status(status):
            p, a = sp.paper, sp.attributes
            first = p.authors[0].split()[0] if p.authors else "[No author]"
            rows.append([
                f"{first} {p.year or ''}", _truncate(p.title), p.journal,
                a.get("design", NOT_STATED), a.get("sample_size", NOT_STATED),
                p.pmid, p.doi or "", f"{sp.status}: {sp.reason} (read: {sp.evidence_basis})",
            ])
    _table(doc, ["First author, year", "Title", "Journal", "Design", "Sample size",
                 "PMID", "DOI", "Our estimate and why"], rows)
    own = c.papers.unclear_without_open_access()
    if own:
        doc.add_paragraph("Retrieve through your own access", style="Heading 3")
        doc.add_paragraph("Unclear from the abstract and not open access: fetch these "
                          "through your library to decide.")
        _bullets(doc, [f"{_format_citation(x.paper)} "
                       f"{'https://doi.org/' + x.paper.doi if x.paper.doi else ''}".strip()
                       for x in own])
    h = c.papers.heterogeneity
    if h is not None and h.distinct_outcomes:
        doc.add_paragraph("Outcome heterogeneity", style="Heading 3")
        verdict = (f"More than {h.threshold}: pooling may be difficult; check outcomes "
                   "before committing to a meta-analysis." if h.flagged
                   else f"Within the {h.threshold} we consider workable.")
        doc.add_paragraph(f"The plausibly eligible papers report {h.distinct_outcomes} "
                          f"distinct primary outcomes. {verdict}")

    # 10. Protocol
    section("Protocol: eligibility criteria and reasoning")
    _bullets(doc, c.criteria.as_text().splitlines())
    for r in c.criteria_reasoning:
        p = doc.add_paragraph()
        p.add_run(f"{r.get('criterion', '')}: ").bold = True
        p.add_run(r.get("reasoning", ""))

    if c.systematic:
        # 11. PROSPERO entry
        section("Pre drafted PROSPERO registration entry")
        if c.prospero:
            _table(doc, ["PROSPERO field", "Draft"],
                   [[k.replace("_", " ").capitalize(), v] for k, v in c.prospero.items()])
        # 12. Registration urgency (verbatim)
        section("Register now")
        for para in REGISTRATION_URGENCY:
            doc.add_paragraph(para)

    # 13. Screening platform
    section("Recommended screening platform")
    doc.add_paragraph(SCREENING_PLATFORM)

    if c.systematic:
        # 14. Risk of bias
        section("Risk of bias tool")
        doc.add_paragraph(c.rob_tool)
    else:
        section("Selecting and appraising sources")
        doc.add_paragraph(f"Appraise key studies with a design-appropriate tool, for "
                          f"example {c.rob_tool}, and say how you selected sources.")

    # 15. Reporting guideline
    section("Reporting guideline")
    _bullets(doc, c.reporting)

    # 16. Aims and objectives
    section("Aims and objectives")
    doc.add_paragraph(c.aims)
    _bullets(doc, c.objectives)

    # 17. Data analysis
    section("Data analysis" if c.systematic else "Synthesis approach")
    doc.add_paragraph(c.analysis)
    if c.skills:
        doc.add_paragraph("Skills you will need", style="Heading 3")
        _bullets(doc, c.skills)

    # 18. Data extraction template
    section("Data extraction template")
    _table(doc, ["Field to extract"], [[x] for x in c.extraction_fields])

    if c.prisma is not None:
        # 19. PRISMA flow diagram
        section("PRISMA flow diagram (preliminary counts)")
        pr = c.prisma
        _table(doc, ["Stage", "Records"], [
            ["Identified through PubMed", f"{pr['identified']:,}"],
            ["Identified through Europe PMC (preprints)", f"{pr.get('preprints_identified', 0):,}"],
            ["Retrieved for our preliminary screen", f"{pr['retrieved']:,}"],
            ["Duplicates removed", f"{pr['duplicates_removed']:,}"],
            ["Titles and abstracts screened", f"{pr['screened']:,}"],
            ["Excluded (our estimate)", f"{pr['excluded_estimate']:,}"],
            ["Full texts to assess", f"{pr['full_text_to_assess']:,}"],
            ["Studies included", "To be determined by your screening"],
        ])

    # 20. Publication structure
    section("Publication structure")
    for heading, subs in c.structure:
        doc.add_paragraph(heading, style="Heading 3")
        _bullets(doc, subs)

    # 21. Grey literature
    section("Grey literature and trial registries")
    _bullets(doc, GREY_LITERATURE)

    # 22. Timeline
    section("Timeline")
    _table(doc, ["Milestone", "Hours (each person)", "Done by week"],
           [[m.phase, f"{m.hours:g}", str(m.week_ends)] for m in c.timeline])
    a = c.timeline_assumptions
    doc.add_paragraph(f"Assumes {a['hours_per_week_each']:g} hours a week each, a team of "
                      f"{a['team']}, about {a['records_to_screen']:,} records and "
                      f"{a['eligible_estimate']} eligible studies. Rates: {a['rates']}.")

    if not c.systematic and c.journals:
        section("Journals to consider")
        doc.add_paragraph("Journals that published the plausibly eligible papers. Check "
                          "each journal's author guidelines and scope before submitting.")
        _bullets(doc, [f"{j} ({k} paper{'s' if k > 1 else ''})" for j, k in c.journals])

    # 23. Common reasons for rejection
    section("Common reasons this type of review is rejected")
    _bullets(doc, rejection_reasons(c.publication_type))

    # 24. Search appendix
    section("Search appendix")
    _table(doc, ["Source", "Query", "Records", "Run"],
           [[x["source"], x["query"], "" if x["count"] is None else f"{x['count']:,}",
             str(x["run_at"])[:10]] for x in c.search_appendix])

    # 25. Collaboration
    section("Find collaborators")
    doc.add_paragraph(COLLABORATION_PROMPT)

    # 26. Disclaimers (verbatim) and disclosure
    section("Disclaimers")
    for para in DISCLAIMERS:
        doc.add_paragraph(para)
    doc.add_paragraph("Acknowledgement you can use", style="Heading 3")
    doc.add_paragraph(ACKNOWLEDGEMENT_WORDING)
    doc.add_paragraph(AUTHOR_AI_NOTE)

    if landscape is not None:
        doc.add_heading("Appendix: research landscape", level=1)
        doc.add_paragraph(
            f"Total literature {landscape.total_literature:,}; systematic reviews "
            f"{landscape.systematic_reviews:,}; guidelines {landscape.guidelines:,}.")

    for style in ("Normal",):
        doc.styles[style].font.size = Pt(11)
    doc.save(str(workspace.guide_path))
    return str(workspace.guide_path)
