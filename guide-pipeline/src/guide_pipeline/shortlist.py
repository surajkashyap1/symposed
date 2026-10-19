"""Pick the question before the guide is written (Samarth, 2026-10-18).

After scoring, a review led run stops with the top three questions, each with a
title, a short project brief, its PICO, why it matters, why it is new and its
scope. A person picks one (`python -m guide_pipeline pick <order> <n>`, which
writes the number into the request folder's CHOICE file) and the run is started
again. Only the picked question goes through the final check, so full texts
are only fetched for a question someone wants. If it fails the final check the
run stops again with the shortlist, the failed option marked, for another pick.
The two questions not picked become the guide's swappable alternatives.

Everything factual (counts, PROSPERO dates) comes from the run's API checks.
The model writes only the PICO, the brief and the rationale.
"""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
from typing import Any, Optional

from .alternatives import Alternative, build
from .guide_content import PLAIN_STYLE
from .plain import tidy

SIZE = 3
CHOICE_FILE = "CHOICE"
FAILED_FILE = "failed_checks.json"
SHORTLIST_MD = "SHORTLIST.md"
SHORTLIST_DOCX = "shortlist.docx"

_S = {"type": "string"}
BRIEF_SCHEMA = {"type": "object", "properties": {"brief": _S}, "required": ["brief"],
                "additionalProperties": False}
_BRIEF_SYSTEM = (
    "You write the project brief for one proposed research question, for the person "
    "deciding which question to take on. In two or three sentences say what the project "
    "is, what the team would actually do, and what they would end up with. Use only the "
    "facts given. Describe the project only: do not comment on the numbers, the request, "
    "other possible questions or yourself, and never write in the first person. "
    + PLAIN_STYLE
)


class AwaitingChoice(Exception):
    """The run needs a person to pick a question. Run again once CHOICE is written."""

    def __init__(self, folder: Path, path: Path, options: list[Alternative], note: str = ""):
        self.folder, self.path, self.options, self.note = folder, path, options, note
        super().__init__(f"pick one of {len(options)} questions: see {path}")


def read_choice(folder: Path) -> Optional[int]:
    try:
        text = (Path(folder) / CHOICE_FILE).read_text(encoding="utf-8").strip()
    except OSError:
        return None
    return int(text) if text.isdigit() else None


def write_choice(folder: Path, number: int) -> Path:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / CHOICE_FILE
    path.write_text(f"{number}\n", encoding="utf-8")
    return path


def failed_checks(folder: Path) -> dict[str, str]:
    """Titles that failed the final check on an earlier pick, with the reason."""
    try:
        return json.loads((Path(folder) / FAILED_FILE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def record_failed(folder: Path, title: str, reason: str) -> None:
    failed = failed_checks(folder)
    failed[title] = reason
    (Path(folder) / FAILED_FILE).write_text(json.dumps(failed, indent=2), encoding="utf-8")
    (Path(folder) / CHOICE_FILE).unlink(missing_ok=True)


def options(scored: list[Any], *, criteria_llm: Any, brief_llm: Any, publication_type: str,
            review_years: int, team: int) -> list[Alternative]:
    """The top questions written up for choosing: PICO, scope, why new, and a brief."""
    out = build([s.assessment for s in scored[:SIZE]], criteria_llm=criteria_llm,
                publication_type=publication_type, review_years=review_years)
    for o in out:
        o.brief = _brief(brief_llm, o, publication_type=publication_type, team=team)
    return out


def _brief(llm: Any, o: Alternative, *, publication_type: str, team: int) -> str:
    pico = "\n".join(f"{k}: {v}" for k, v in o.pico)
    user = (f'Proposed {publication_type}: "{o.title}"\n{pico}\nWhy it matters: {o.rationale}\n'
            f"Size: about {o.eligible_estimate} studies likely to fit, from about "
            f"{o.records_to_screen} search results to screen.\n"
            f"The team: {team} {'person' if team == 1 else 'people'}.")
    try:
        return tidy(str(llm.complete_json(_BRIEF_SYSTEM, user, schema=BRIEF_SCHEMA)["brief"]))
    except Exception:  # noqa: BLE001  a missing brief never blocks the shortlist
        return ""


def alternatives_after(opts: list[Alternative], chosen: int, *, failed: set[str]) -> list[Alternative]:
    """The options not picked (and not failed), renumbered from 1, as the guide's alternatives."""
    rest = [o for o in opts if o.number != chosen and o.title not in failed]
    return [replace(o, number=i) for i, o in enumerate(rest, 1)]


def considered(a: Any, *, number: Optional[int], topic: str, novelty: Any = None) -> dict:
    """One question from the search, for the full list a person approves from."""
    c, v = a.candidate, a.verdict
    return {"number": number, "title": c.title, "topic": topic, "axis": c.axis,
            "rationale": tidy(c.rationale), "passed_gates": v.passed,
            "gate_reasons": [tidy(r) for r in v.reasons if r],
            "eligible": a.counts.eligible_studies, "records": a.counts.records_to_screen,
            "novelty": novelty.as_dict() if novelty is not None else {}}


def _considered_md(others: list[dict]) -> list[str]:
    if not others:
        return []
    lines = ["", "# Every question considered", "",
             "All the questions the search produced, by topic, whatever their rating. Any "
             "numbered one can be picked with the same command. Questions without a number "
             "did not pass the basic checks (too few studies, already reviewed or registered) "
             "and cannot be written up as they stand.", ""]
    by_topic: dict[str, list[dict]] = {}
    for o in others:
        by_topic.setdefault(o["topic"] or "Other", []).append(o)
    for name, items in by_topic.items():
        lines += [f"## Topic: {name}", ""]
        for o in sorted(items, key=lambda x: (x["number"] is None, x["number"] or 0)):
            head = (f"### Option {o['number']}: {o['title']}" if o["number"]
                    else f"### (not pickable) {o['title']}")
            lines += [head, ""]
            nv = o.get("novelty") or {}
            if nv:
                lines += [f"**Novelty check: {VERDICT_LABEL.get(nv['verdict'], nv['verdict'])}.** "
                          f"{nv.get('editor_view', '')}", ""]
                closest = next((e for e in nv.get("papers", [])
                                if e["pmid"] == nv.get("closest_pmid")), None)
                if closest:
                    lines += [f"Closest paper: [{closest['title']}](https://pubmed.ncbi.nlm.nih.gov/"
                              f"{closest['pmid']}/) ({closest['journal']}, {closest.get('year') or ''})", ""]
            elif o["passed_gates"]:
                lines += ["**Novelty check:** not run (the limit per topic was reached).", ""]
            if not o["passed_gates"]:
                lines += [f"**Did not pass the basic checks:** {' '.join(o['gate_reasons'])}", ""]
            lines += [f"**Why it was proposed.** {o['rationale']}", "",
                      f"About {o['eligible']} studies likely to fit, from {o['records']} search "
                      "results.", ""]
    return lines


def write(folder: Path, opts: list[Alternative], *, topic: str, key: str,
          failed: dict[str, str], note: str = "", others: list[dict] = ()) -> Path:
    """SHORTLIST.md (to read or paste) and shortlist.docx (to send)."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    lines = [f"# Pick a question: {topic}", ""]
    if note:
        lines += [f"**{note}**", ""]
    lines += ["Read the options, pick one, then run:", "",
              f"    python -m guide_pipeline pick {key} <number>", "",
              "and run the batch again. The guide is written for the question picked, and "
              "the other two become its alternatives.", ""]
    for o in opts:
        lines += [f"## Option {o.number}: {o.title}", ""]
        if o.title in failed:
            lines += [f"**Failed the final check, pick another:** {failed[o.title]}", ""]
        if o.brief:
            lines += [f"**Project brief.** {o.brief}", ""]
        lines += ["**PICO**", ""] + [f"- {k}: {v}" for k, v in o.pico] + [""]
        lines += [f"**Why it matters.** {o.rationale}", "", f"**Why it is new.** {o.why_novel}", ""]
        lines += _novelty_md(o.novelty)
        lines += ["**Scope**", ""] + [f"- {k}: {v}" for k, v in o.scope] + [""]
        lines += [f"Best search ({o.search_count} results in PubMed): `{o.search_query}`", ""]
    others = [{**o, "rationale": f"Failed the final check: {failed[o['title']]} {o['rationale']}"}
              if o["title"] in failed else o for o in others]
    lines += _considered_md(others)
    path = folder / SHORTLIST_MD
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    _write_docx(folder / SHORTLIST_DOCX, opts, topic=topic, failed=failed, note=note,
                others=others)
    return path


VERDICT_LABEL = {"novel": "New, as far as the closest papers show",
                 "uncertain": "Uncertain: a person must check the closest papers",
                 "not novel": "Not new", "dead": "New but unlikely to give a useful answer"}


def _papers(novelty: dict) -> list[dict]:
    order = {"same question": 0, "large": 1, "partial": 2, "background only": 3, "": 4}
    closest = novelty.get("closest_pmid", "")
    return sorted(novelty.get("papers", []),
                  key=lambda e: (e["pmid"] != closest, order.get(e.get("overlap", ""), 4)))


def _novelty_md(novelty: dict) -> list[str]:
    if not novelty:
        return []
    lines = [f"**Novelty check (editor's view): {VERDICT_LABEL.get(novelty['verdict'], novelty['verdict'])}.** "
             f"{novelty.get('editor_view', '')}", ""]
    for label, key in (("What it would add", "adds"), ("Worth answering whatever it finds?",
                                                       "worth_answering"),
                       ("To make it stronger", "suggested_change")):
        if novelty.get(key):
            lines += [f"**{label}** {novelty[key]}", ""]
    papers = _papers(novelty)
    if papers:
        lines += ["**Closest existing papers** (check these yourself)", ""]
        for e in papers:
            tag = " **(closest)**" if e["pmid"] == novelty.get("closest_pmid") else ""
            overlap = f" Overlap: {e['overlap']}." if e.get("overlap") else ""
            why = f" {e['why']}" if e.get("why") else ""
            lines.append(f"- [{e['title']}](https://pubmed.ncbi.nlm.nih.gov/{e['pmid']}/) "
                         f"({e['journal']}, {e.get('year') or ''}, {e['kind']}){tag}.{overlap}{why}")
        lines.append("")
    return lines


def _write_docx(path: Path, opts: list[Alternative], *, topic: str, failed: dict[str, str],
                note: str, others: list[dict] = ()) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(f"Pick a question: {topic}", 0)
    if note:
        doc.add_paragraph().add_run(note).bold = True
    doc.add_paragraph("Read the options below and reply with the number you want. We then "
                      "check the chosen question once more against the closest existing "
                      "reviews and write its guide. The other two become your alternatives.")
    for o in opts:
        doc.add_heading(f"Option {o.number}: {o.title}", 1)
        if o.title in failed:
            p = doc.add_paragraph()
            p.add_run("Failed the final check, pick another: ").bold = True
            p.add_run(failed[o.title])
        for label, text in (("Project brief", o.brief), ("Why it matters", o.rationale),
                            ("Why it is new", o.why_novel)):
            if text:
                p = doc.add_paragraph()
                p.add_run(f"{label}. ").bold = True
                p.add_run(text)
        if o.novelty:
            nv = o.novelty
            p = doc.add_paragraph()
            p.add_run("Novelty check (editor's view): "
                      f"{VERDICT_LABEL.get(nv['verdict'], nv['verdict'])}. ").bold = True
            p.add_run(nv.get("editor_view", ""))
            for label, key in (("What it would add. ", "adds"),
                               ("Worth answering whatever it finds? ", "worth_answering"),
                               ("To make it stronger. ", "suggested_change")):
                if nv.get(key):
                    q = doc.add_paragraph()
                    q.add_run(label).bold = True
                    q.add_run(nv[key])
            papers = _papers(nv)
            if papers:
                doc.add_heading("Closest existing papers (check these yourself)", 2)
                for e in papers:
                    q = doc.add_paragraph(style="List Bullet")
                    q.add_run(e["title"]).bold = e["pmid"] == nv.get("closest_pmid")
                    q.add_run(f" ({e['journal']}, {e.get('year') or ''}, {e['kind']}, "
                              f"PMID {e['pmid']}, https://pubmed.ncbi.nlm.nih.gov/{e['pmid']}/)")
                    if e.get("overlap"):
                        q.add_run(f" Overlap: {e['overlap']}. {e.get('why', '')}")
        for heading, rows in (("PICO", o.pico), ("Scope", o.scope)):
            if not rows:
                continue
            doc.add_heading(heading, 2)
            table = doc.add_table(rows=0, cols=2)
            table.style = "Light Grid Accent 1"
            for k, v in rows:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = k, v
    if others:
        doc.add_heading("Every question considered", 1)
        doc.add_paragraph("All the questions the search produced, by topic, whatever their "
                          "rating. Reply with the number of any numbered one to approve it. "
                          "Unnumbered ones did not pass the basic checks.")
        topic_now = None
        for o in sorted(others, key=lambda x: (x["topic"], x["number"] is None, x["number"] or 0)):
            if o["topic"] != topic_now:
                topic_now = o["topic"]
                doc.add_heading(f"Topic: {topic_now or 'Other'}", 2)
            label = f"Option {o['number']}: " if o["number"] else "(not pickable) "
            doc.add_heading(label + o["title"], 3)
            nv = o.get("novelty") or {}
            if nv:
                q = doc.add_paragraph()
                q.add_run(f"Novelty check: {VERDICT_LABEL.get(nv['verdict'], nv['verdict'])}. ").bold = True
                q.add_run(nv.get("editor_view", ""))
            if not o["passed_gates"]:
                q = doc.add_paragraph()
                q.add_run("Did not pass the basic checks: ").bold = True
                q.add_run(" ".join(o["gate_reasons"]))
            q = doc.add_paragraph()
            q.add_run("Why it was proposed. ").bold = True
            q.add_run(o["rationale"])
    doc.save(str(path))
