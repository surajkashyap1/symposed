from datetime import date, datetime, timezone
from types import SimpleNamespace

from docx import Document

from guide_pipeline.candidates import Candidate, CandidateAssessment, RequestPreferences, SearchRun
from guide_pipeline.gates import GateVerdict, TriageCounts
from guide_pipeline.guide import _format_citation, build_guide
from guide_pipeline.guide_content import (
    DISCLAIMERS,
    PROSPERO_SCHEMA,
    PROTOCOL_SCHEMA,
    REGISTRATION_URGENCY,
    STRATEGY_SCHEMA,
    _clean_citations,
    assemble,
    journal_shortlist,
    reporting_guideline,
    risk_of_bias_tool,
    timeline,
)
from guide_pipeline.prospero import ProsperoCheck, ProsperoMatch
from guide_pipeline.retrieval import RetrievalResult
from guide_pipeline.screening import (
    ABSTRACT_BASIS,
    LIKELY_ELIGIBLE,
    LIKELY_INELIGIBLE,
    UNCLEAR,
    Criteria,
    HeterogeneityCheck,
    ScreenedPaper,
    ScreeningResult,
)
from guide_pipeline.sources.pubmed import Paper
from guide_pipeline.workspace import create_workspace


def _paper(pmid, title, authors=("Smith J",), year=2021, journal="J Test", doi="10.1/x"):
    return Paper(pmid, title, "abstract", tuple(authors), journal, year, doi)


CRIT = Criteria("ICU adults", "Vitamin D", "Placebo", "Mortality", "Randomised controlled trials",
                exclusion=("Children",))


def _screening():
    papers = [
        ScreenedPaper(_paper("111", "First trial"), {"design": "RCT", "sample_size": "120"},
                      LIKELY_ELIGIBLE, "Matches population and intervention.",
                      "full text (methods and results)", "PMC111"),
        ScreenedPaper(_paper("222", "Unclear one", doi="10.1/unclear"), {}, UNCLEAR,
                      "Outcome not stated in the abstract.", ABSTRACT_BASIS),
        ScreenedPaper(_paper("333", "Paediatric study"), {}, LIKELY_INELIGIBLE,
                      "Population is children.", ABSTRACT_BASIS),
    ]
    return ScreeningResult(CRIT, papers, HeterogeneityCheck({"death": ["111"]}, 5),
                           full_text_screened=1,
                           counts={LIKELY_ELIGIBLE: 1, UNCLEAR: 1, LIKELY_INELIGIBLE: 1})


class FakePubMed:
    def search_pmids(self, query, retmax=400):
        return ["900", "901"]

    def titles(self, pmids):
        return [("900", "Vitamin D in the ICU: a meta-analysis"), ("901", "Unrelated cardiology")]

    def count(self, query, **kw):
        return 0 if "broken" in query else 321

    def mesh_terms(self, term, max_terms=5):
        return [SimpleNamespace(name="Vitamin D")] if term == "Vitamin D" else []


class GuideLLM:
    def __init__(self, recommend=0):
        self.calls, self.recommend = [], recommend

    def complete_json(self, system, user, schema=None, cache_system=False):
        self.calls.append(schema)
        if schema is STRATEGY_SCHEMA:
            return {"strategies": [
                {"name": "Sensitive", "pubmed_query": '"Vitamin D"[mh] AND "Madeup Term"[mh]',
                 "purpose": "Find everything"},
                {"name": "Precise", "pubmed_query": "broken query", "purpose": "Narrow"}],
                "recommended": self.recommend, "recommendation_reason": "Best recall",
                "embase_emtree": "'vitamin d'/exp", "cochrane_central": "[mh 'Vitamin D']"}
        if schema is PROTOCOL_SCHEMA:
            return {"rationale": "Evidence is mixed [1] and new [2] but see [99].",
                    "similar_work_differences": [{"ref": 1, "difference": "Ours is adults only."},
                                                 {"ref": 42, "difference": "invented"}],
                    "criteria_reasoning": [{"criterion": "Population", "reasoning": "ICU only"}],
                    "aims": "To assess vitamin D.", "objectives": ["Pool mortality"],
                    "analysis": "Random effects meta-analysis.", "skills_needed": ["RevMan"]}
        if schema is PROSPERO_SCHEMA:
            return {k: f"draft {k}" for k in PROSPERO_SCHEMA["properties"]}
        raise AssertionError(schema)


def _assessment():
    m = ProsperoMatch("CRD9", "A registered protocol", date(2025, 1, 1), 0.4)
    return CandidateAssessment(
        candidate=Candidate("Vitamin D in ICU", "Discordance", "vitamin d[tiab]", rationale="gap"),
        counts=TriageCounts(300, 40, 0, 0),
        verdict=GateVerdict("pass", []),
        prospero=ProsperoCheck("t", "t", date(2026, 9, 28), date(2026, 9, 28),
                               datetime(2026, 9, 28, tzinfo=timezone.utc), "clear", [m]),
        searches=[SearchRun("PubMed", "vitamin d[tiab]", 300, "2026-09-28T10:00:00")],
    )


def _content(publication_type="systematic review", llm=None):
    llm = llm or GuideLLM()
    return assemble(
        title="Vitamin D in ICU", publication_type=publication_type, criteria=CRIT,
        assessment=_assessment(),
        retrieval=RetrievalResult("vitamin d[tiab] NOT review", [], identified=300, fetched=3),
        screening=_screening(),
        preferences=RequestPreferences(collaborators=3, context="The user is an SHO."),
        proforma={"hoursPerWeek": "3 to 5"}, search_date="2026-09-28",
        pubmed=FakePubMed(),
        llms={"search_strategy": llm, "protocol_drafting": llm, "prospero_form": llm},
    ), llm


def _text(path):
    doc = Document(path)
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def test_format_citation_variants():
    many = _paper("1", "T", authors=("A A", "B B", "C C", "D D"))
    assert _format_citation(many).startswith("A A, B B, C C, et al. (2021). T.")
    no_doi = _format_citation(_paper("2", "T2", authors=(), doi=None))
    assert "PMID: 2." in no_doi and "DOI:" not in no_doi


def test_strategies_get_real_counts_and_mesh_is_checked():
    c, _ = _content()
    sensitive, precise = c.strategies["strategies"]
    assert sensitive.count == 321 and sensitive.unverified_mesh == ["Madeup Term"]
    assert precise.count == 0
    assert any("Madeup Term" in w for w in c.warnings)
    assert any("returned 0 records" in w for w in c.warnings)


def test_invented_citations_and_references_are_removed():
    c, _ = _content()
    assert "[99]" not in c.rationale and "[1]" in c.rationale and "[2]" in c.rationale
    assert c.differences == {1: "Ours is adults only."}  # ref 42 did not exist
    assert _clean_citations("see [0] and [3]", 3) == "see  and [3]"


def test_similar_work_combines_reviews_and_prospero_ranked():
    c, _ = _content()
    titles = [s["title"] for s in c.similar]
    assert titles[0] == "Vitamin D in the ICU: a meta-analysis"
    assert any(t.startswith("PROSPERO CRD9") for t in titles)


def test_systematic_guide_has_every_spec_section(tmp_path):
    c, llm = _content()
    ws = create_workspace(base=tmp_path, request_id="sr")
    text = _text(build_guide(ws, c))
    for heading in ("Novelty statement", "PICO", "Rationale", "Similar existing work",
                    "Feasibility", "Search strategies", "Translated syntax",
                    "Suggested papers", "Protocol", "PROSPERO registration entry",
                    "Register now", "screening platform", "Risk of bias tool",
                    "Reporting guideline", "Aims and objectives", "Data analysis",
                    "Data extraction template", "PRISMA flow diagram",
                    "Publication structure", "Grey literature", "Timeline",
                    "rejected", "Search appendix", "Find collaborators", "Disclaimers"):
        assert heading in text, heading
    assert ("This title was checked against PROSPERO and the published literature on "
            "28/09/2026") in text
    for para in (*REGISTRATION_URGENCY, *DISCLAIMERS):
        assert para in text
    assert "Methodological support for this study was provided by Symposed." in text
    assert "321" in text  # a real strategy count
    assert "https://doi.org/10.1/unclear" in text  # own-access list
    assert "draft review_question" in text  # PROSPERO entry
    assert len(llm.calls) == 3


def test_literature_review_skips_sr_only_sections(tmp_path):
    c, llm = _content("literature review")
    ws = create_workspace(base=tmp_path, request_id="lit")
    text = _text(build_guide(ws, c))
    assert "PROSPERO registration entry" not in text and "PRISMA flow diagram" not in text
    assert "Register now" not in text
    assert "Journals to consider" in text and "J Test (2 papers)" in text
    assert "Selecting and appraising sources" in text
    assert PROSPERO_SCHEMA not in llm.calls  # no PROSPERO call for a literature review


def test_rule_based_choices():
    assert risk_of_bias_tool("Randomised controlled trials") == "RoB 2 (randomised trials)"
    assert "ROBINS-I" in risk_of_bias_tool("RCTs and cohort studies (randomised or not)")
    assert risk_of_bias_tool("studies", "Diagnostic accuracy").startswith("QUADAS-2")
    assert risk_of_bias_tool("prediction model studies").startswith("PROBAST")
    assert reporting_guideline("systematic review")[0] == "PRISMA 2020"
    assert reporting_guideline("scoping or narrative review")[0].startswith("PRISMA-ScR")


def test_timeline_scales_with_hours_and_team():
    slow, _ = timeline(records=400, eligible=30, team=2, hours_label="Under 3",
                       publication_type="systematic review")
    fast, a = timeline(records=400, eligible=30, team=4, hours_label="Over 10",
                       publication_type="systematic review")
    assert slow[-1].week_ends > fast[-1].week_ends
    assert [m.week_ends for m in slow] == sorted(m.week_ends for m in slow)
    assert a["team"] == 4


def test_journal_shortlist_counts_only_plausible_papers():
    assert journal_shortlist(_screening()) == [("J Test", 2)]


def test_a_zero_record_strategy_is_never_recommended():
    c, _ = _content(llm=GuideLLM(recommend=1))  # the model picks the broken one
    assert c.strategies["recommended"] == 0
    assert "321" in c.strategies["recommendation_reason"]
    assert any("switched to 'Sensitive'" in w for w in c.warnings)


def test_screening_volume_uses_the_triage_count():
    c, _ = _content()
    assert c.feasibility["records_to_screen"] == 300  # triage count, not the 3 fetched
    assert c.timeline_assumptions["records_to_screen"] == 300
