"""Open-access full text (JATS XML) reduced to what eligibility depends on.

Screening needs who was studied, what was done and what was measured — the
methods and results. References, funding, declarations and the like are dropped
so each paper costs fewer tokens. Many articles don't label their sections, so
sections are chosen by title; when none match, the whole body is kept minus the
administrative sections. Anything cut for length is stated in the evidence basis.
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass

_KEEP = re.compile(
    r"method|material|result|participant|patient|population|subject|design|"
    r"intervention|outcome|eligib|setting|procedure|protocol|trial|finding",
    re.IGNORECASE,
)
_DROP = re.compile(
    r"reference|acknowledg|funding|financ|declaration|competing|conflict|"
    r"contribution|abbreviation|footnote|supplementary|availability|contributor|"
    r"associated data|ethic|consent for publication|author",
    re.IGNORECASE,
)
_DROP_TYPES = {"ref-list", "ack", "glossary", "fn-group", "contrib-info", "associated-data",
               "supplementary-material", "COI-statement"}


@dataclass(frozen=True)
class FullText:
    text: str
    sections: tuple[str, ...]  # section titles used
    truncated: bool

    @property
    def basis(self) -> str:
        """How the screening evidence basis describes this text."""
        which = (
            "methods and results"
            if any(_KEEP.search(t) for t in self.sections)
            else "body"
        )
        return f"full text ({which}{', truncated' if self.truncated else ''})"


def _text(node: ET.Element) -> str:
    return " ".join(" ".join(node.itertext()).split())


def extract_screening_text(xml: str, *, max_chars: int = 60_000) -> FullText | None:
    """Title + abstract + eligibility-relevant sections, or None if unusable."""
    try:
        root = ET.fromstring(xml)
    except ET.ParseError:
        return None
    article = root if root.tag == "article" else root.find(".//article")
    if article is None:
        return None
    body = article.find("body")
    if body is None:
        return None  # no body: nothing beyond the abstract

    front = article.find("front")
    title = _text(front.find(".//article-title")) if front is not None and front.find(
        ".//article-title") is not None else ""
    abstract = (
        _text(front.find(".//abstract"))
        if front is not None and front.find(".//abstract") is not None
        else ""
    )

    sections = [
        s for s in body.findall("sec")
        if s.get("sec-type") not in _DROP_TYPES and not _DROP.search(s.findtext("title") or "")
    ]
    chosen = [s for s in sections if _KEEP.search(s.findtext("title") or "")] or sections
    if not chosen:
        return None
    parts = [f"TITLE: {title}", f"ABSTRACT: {abstract}"]
    parts += [f"SECTION: {(s.findtext('title') or '').strip()}\n{_text(s)}" for s in chosen]
    text = "\n\n".join(parts)
    truncated = len(text) > max_chars
    if truncated:
        text = text[:max_chars]
    return FullText(
        text=text,
        sections=tuple((s.findtext("title") or "").strip() for s in chosen),
        truncated=truncated,
    )


# -- whole reviews (discovery) ----------------------------------------------------
# A review's gaps and limitations live in its discussion, so for reading a review
# everything is kept except the administrative sections.
def extract_review_text(xml: str, *, max_chars: int = 120_000) -> FullText | None:
    """Title + abstract + every body section except references and admin."""
    try:
        root = ET.fromstring(xml)
    except ET.ParseError:
        return None
    article = root if root.tag == "article" else root.find(".//article")
    body = article.find("body") if article is not None else None
    if body is None:
        return None
    front = article.find("front")
    title_node = front.find(".//article-title") if front is not None else None
    abstract_node = front.find(".//abstract") if front is not None else None
    sections = [
        s for s in body.findall("sec")
        if s.get("sec-type") not in _DROP_TYPES and not _DROP.search(s.findtext("title") or "")
    ]
    if not sections:
        return None
    parts = [f"TITLE: {_text(title_node) if title_node is not None else ''}",
             f"ABSTRACT: {_text(abstract_node) if abstract_node is not None else ''}"]
    parts += [f"SECTION: {(s.findtext('title') or '').strip()}\n{_text(s)}" for s in sections]
    text = "\n\n".join(parts)
    truncated = len(text) > max_chars
    return FullText(text=text[:max_chars],
                    sections=tuple((s.findtext("title") or "").strip() for s in sections),
                    truncated=truncated)


LOCAL_SUFFIXES = (".docx", ".pdf", ".txt", ".md")


def read_local_text(path) -> str:
    """Text of a full text file supplied by hand: Word, PDF, plain text or Markdown.

    Most will be text copied from the publisher's page into a Word document.
    """
    from pathlib import Path

    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(p))
        paras = [x.text for x in doc.paragraphs]
        cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        return "\n".join(x for x in paras + cells if x.strip())
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages)
    if suffix in (".txt", ".md"):
        return p.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported full text file: {p.name}")
