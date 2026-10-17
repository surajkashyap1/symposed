import re
import zipfile
from html import unescape

import openpyxl
import pytest
from docx import Document

from guide_pipeline import diagrams
from guide_pipeline.companions import write_all
from guide_pipeline.deliver import build_outputs, rerender
from guide_pipeline.guide_model import likely_publication_type, paper_rows, step_timeline, steps
from guide_pipeline.guide_pdf import render_html
from guide_pipeline.plain import glossary_for, tidy
from guide_pipeline.rob_tools import choose_tools
from guide_pipeline.screening import (
    ABSTRACT_BASIS,
    LIKELY_ELIGIBLE,
    LIKELY_INELIGIBLE,
    UNCLEAR,
    ScreenedPaper,
    ScreeningResult,
)
from guide_pipeline.sources.pubmed import Paper
from guide_pipeline.workspace import create_workspace
from sample_model import sample_model
from test_guide import CRIT


def test_tidy_removes_semicolons_and_dashes_but_keeps_hyphens():
    out = tidy("Trials differ; some are small — others large. A meta-analysis of 2015–2020 data")
    assert ";" not in out and "—" not in out and "–" not in out
    assert "Trials differ. Some are small, others large." in out
    assert "meta-analysis" in out and "2015 to 2020" in out
    assert tidy("one - two") == "one, two"


def test_glossary_picks_terms_used_plus_project_terms():
    g = glossary_for("Use RoB 2 and PubMed with MeSH terms.",
                     [{"term": "Hall technique", "definition": "A crown; no drilling."}])
    terms = [x["term"] for x in g]
    assert {"RoB 2", "PubMed", "MeSH", "Hall technique"} <= set(terms)
    assert "PROBAST" not in terms
    assert terms == sorted(terms, key=str.lower)
    assert ";" not in next(x for x in g if x["term"] == "Hall technique")["definition"]


def test_likely_publication_type_prefers_specific_pubmed_types():
    assert likely_publication_type(["Journal Article", "Randomized Controlled Trial"], "") == \
        ("Randomised controlled trial", "PubMed")
    assert likely_publication_type(["Journal Article", "Review", "Systematic Review"], "") == \
        ("Systematic review", "PubMed")
    assert likely_publication_type(["Journal Article"], "retrospective cohort") == \
        ("Retrospective cohort", "our estimate")
    assert likely_publication_type([], "not stated")[1] == "our estimate"
    assert likely_publication_type([], "", preprint=True)[0].startswith("Preprint")


def test_paper_rows_list_only_papers_worth_reading_with_no_rating():
    p = lambda pmid, types=(): Paper(pmid, f"T{pmid}", "abs", ("Smith J",), "J", 2021, None,
                                     publication_types=types)
    sr = ScreeningResult(CRIT, [
        ScreenedPaper(p("1", ("Randomized Controlled Trial",)), {"design": "RCT"},
                      LIKELY_ELIGIBLE, "ok", ABSTRACT_BASIS),
        ScreenedPaper(p("2"), {"design": "cohort"}, UNCLEAR, "?", ABSTRACT_BASIS),
        ScreenedPaper(p("3"), {}, LIKELY_INELIGIBLE, "no", ABSTRACT_BASIS),
    ], None)
    rows = paper_rows(sr)
    assert [r["pmid"] for r in rows] == ["1", "2"]
    assert all("status" not in r and "reason" not in r for r in rows)
    assert rows[0]["publication_type"] == "Randomised controlled trial"
    assert rows[1]["publication_type_basis"] == "our estimate"


def test_step_timeline_uses_the_step_names_and_scales_with_hours():
    slow, _ = step_timeline(systematic=True, records=400, eligible=30, team=2,
                            hours_label="Under 3")
    fast, a = step_timeline(systematic=True, records=400, eligible=30, team=4,
                            hours_label="Over 10")
    assert [t["name"] for t in slow] == [s["name"] for s in steps(True)]
    assert slow[-1]["end_week"] > fast[-1]["end_week"]
    assert all(t["start_week"] <= t["end_week"] for t in slow)
    assert [t["start_week"] for t in slow] == sorted(t["start_week"] for t in slow)
    assert a["total_weeks"] == fast[-1]["end_week"]


def test_tools_match_the_designs():
    assert [t.name for t in choose_tools("Systematic reviews of the Hall technique")] == ["ROBIS"]
    assert [t.name for t in choose_tools("prediction model studies")] == ["PROBAST"]
    assert "CHARMS checklist" in [c[0] for c in choose_tools("prediction model studies")[0].companions]
    assert [t.name for t in choose_tools("RCTs", systematic=False)] == ["CASP checklists", "RoB 2"]


def _visible_prose(html: str) -> str:
    html = re.sub(r"<(code|svg|style)[^>]*>.*?</\1>", " ", html, flags=re.S)
    html = re.sub(r'<span class="query">.*?</span>', " ", html, flags=re.S)
    return unescape(re.sub(r"<[^>]+>", " ", html))


@pytest.mark.parametrize("systematic", [True, False])
def test_guide_text_follows_house_style(systematic):
    prose = _visible_prose(render_html(sample_model(systematic)))
    assert ";" not in prose
    assert "—" not in prose and "–" not in prose
    assert not re.search(r"\s-\s", prose)


def test_papers_intro_names_the_numbered_search_and_disclaimers_are_red():
    m = sample_model()
    html = render_html(m)
    assert "These are the papers we found using <strong>search strategy 1</strong>" in html
    assert html.count('class="box box-warn"') >= 7
    assert "comes from your draft protocol" in html
    assert "Likely publication type" in html and "likely ineligible" not in html.lower()
    for st in m["steps"]:
        assert f"Step {st['number']} of 9" in html


def test_diagrams_escape_text_and_grow_with_content():
    svg = diagrams.decision_tree([{"question": "Is it <an> RCT & randomised?", "no_label": "x"}])
    assert "&lt;an&gt;" in svg and "&amp;" in svg and svg.startswith("<svg")
    short = diagrams.decision_tree([{"question": "Q?", "no_label": "x"}])
    long = diagrams.decision_tree([{"question": "Q? " * 60, "no_label": "x"}])
    h = lambda s: float(re.search(r'height="([\d.]+)"', s).group(1))
    assert h(long) > h(short)
    assert "Wk 1" in diagrams.gantt([{"number": 1, "name": "A", "start_week": 1, "end_week": 3}])


def test_companions_are_bespoke_and_editable(tmp_path):
    m = sample_model()
    out = write_all(m, tmp_path)
    assert set(out) == {"tracker", "protocol", "prospero"}
    wb = openpyxl.load_workbook(out["tracker"])
    assert wb.sheetnames[:5] == ["How to use", "Screening", "Data extraction", "Risk of bias",
                                 "Search log"]
    sh = wb["Screening"]
    assert sh.max_row == len(m["papers"]) + 1
    headers = [c.value for c in sh[1]]
    for col in ("PMID or ID", "Abstract", "Likely publication type", "Screener 1 decision",
                "Reason for exclusion"):
        assert col in headers
    lists = wb["Lists"]
    values = [c.value for col in lists.iter_cols() for c in col]
    assert "Further review needed" in values and "Wrong population" in values
    assert any("RoB 2: Bias" in (c.value or "") for c in wb["Risk of bias"][1])
    assert [c.value for c in wb["Data extraction"][1]][2] == m["extraction_fields"][0]
    doc = Document(out["protocol"])
    text = "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert m["title"] in text and "starting point" in text
    (tmp_path / "lit").mkdir()
    lit = write_all(sample_model(systematic=False), tmp_path / "lit")
    assert "prospero" not in lit


def test_build_outputs_writes_pack_and_rerenders(tmp_path):
    pytest.importorskip("weasyprint")
    ws = create_workspace(base=tmp_path, request_id="r")
    out = build_outputs(ws, sample_model(), screenshots=False, say=lambda _m: None)
    assert out["pdf"].read_bytes()[:4] == b"%PDF"
    assert sorted(zipfile.ZipFile(out["pack"]).namelist()) == [
        "guide.pdf", "prospero_draft.docx", "protocol.docx", "tracker.xlsx"]
    again = rerender(ws.dir, say=lambda _m: None)
    assert again["pdf"].exists() and ws.model_path.exists()


def test_long_extracted_values_are_shortened_for_narrow_columns():
    from guide_pipeline.guide_model import shorten
    assert shorten("5537 patients (2575 MITT and 2962 SITT users)", 28) == "5537 patients"
    assert shorten("Retrospective cohort study (new-user, active comparator) using CPRD", 45) == \
        "Retrospective cohort study"
    long = shorten("4659 FF/UMEC/VI initiators and 9845 MITT initiators for the population", 28)
    assert len(long) <= 29 and long.endswith("…")
    # a vague PubMed label loses to a specific design read from the abstract
    assert likely_publication_type(["Comparative Study"], "retrospective cohort study") == \
        ("Retrospective cohort study", "our estimate")
    assert likely_publication_type(["Comparative Study"], "") == ("Comparative study", "PubMed")
